#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del Dashboard de Seguimiento de Creditos por Venta - World Salud.

Lee la planilla "Seguimiento Facturacion Clientes 2026" (Google Sheet) y produce
un archivo HTML con la fecha del corte en el nombre, por ejemplo:

    DSH_SEGUIMIENTO_CLIENTES_WS_11062026.html

USO
---
1) En Google Sheets:  Archivo > Descargar > Microsoft Excel (.xlsx)
2) Ejecutar:

       python generar_dashboard.py  Seguimiento_Facturacion.xlsx

   Opcional, fijar la fecha del corte (por defecto usa la fecha de hoy):

       python generar_dashboard.py  archivo.xlsx  --fecha 08/06/2026

Si no se pasa un archivo, intenta descargar la planilla directamente desde
Google (solo funciona si la planilla esta compartida "cualquiera con el enlace").

REQUISITOS:  pip install openpyxl
"""
import sys, os, json, datetime, urllib.request

SHEET_ID  = "17RdCgz50wF5pSqS0GcAKRv4wblUrIBetP-DyF0tEako"
PLANTILLA = "plantilla_dashboard.html"
MES_ABR   = ['Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalalo con:  pip install openpyxl")


# ---------------------------------------------------------------- utilidades
def es_fecha(v):
    return isinstance(v, (datetime.datetime, datetime.date))

def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0

def buscar_fila(ws, pred, hasta=None):
    for r in range(1, (hasta or ws.max_row) + 1):
        if pred(ws, r):
            return r
    return None


# ---------------------------------------------------------------- parsers
def parse_facturacion(ws):
    """Bloque FACTURACION del tab 'Tableros'. Devuelve (meses, {cliente:[...]})."""
    hdr = buscar_fila(ws, lambda w, r: str(w.cell(r, 1).value or "").strip() == "Tipo de Servicio")
    # columnas-mes = celdas fecha en la fila de encabezado
    cols = [(c, ws.cell(hdr, c).value.month) for c in range(3, ws.max_column + 1) if es_fecha(ws.cell(hdr, c).value)]
    meses = [MES_ABR[m - 1] for _, m in cols]
    # columna "Suma total" (autoridad del total por cliente)
    total_col = next((c for c in range(3, ws.max_column + 1)
                      if str(ws.cell(hdr, c).value or "").strip().startswith("Suma total")), None)
    data = {}
    for r in range(hdr + 1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        b = str(ws.cell(r, 2).value or "").strip()
        if a.startswith("Suma total"):
            break
        if not b:                       # filas "Total Farmacia", etc.
            continue
        vals = [round(num(ws.cell(r, c).value), 2) for c, _ in cols]
        # plegar al ultimo mes cualquier columna extra sin encabezado de mes
        if total_col is not None and vals:
            resid = round(num(ws.cell(r, total_col).value) - sum(vals), 2)
            if abs(resid) > 1:
                vals[-1] = round(vals[-1] + resid, 2)
        data[b] = vals
    return meses, data

def parse_ingresos(ws, meses):
    """Bloque INGRESOS. Alinea cada mes al indice de 'meses' (superset)."""
    hdr = buscar_fila(ws, lambda w, r: str(w.cell(r, 1).value or "").strip() == "Cliente"
                      and str(w.cell(r, 2).value or "").strip().startswith("N"))
    cols = [(c, ws.cell(hdr, c).value.month) for c in range(3, ws.max_column + 1) if es_fecha(ws.cell(hdr, c).value)]
    idx = {MES_ABR[m - 1]: i for i, m in enumerate([mm for _, mm in
            [(c, ws.cell(hdr, c).value.month) for c in range(3, ws.max_column + 1) if es_fecha(ws.cell(hdr, c).value)]])}
    data = {}
    for r in range(hdr + 1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        if a.startswith("Suma total") or not a:
            if a.startswith("Suma total"):
                break
            continue
        cli = a[len("Total "):] if a.startswith("Total ") else a
        arr = [0.0] * len(meses)
        for c, m in cols:
            lab = MES_ABR[m - 1]
            if lab in meses:
                arr[meses.index(lab)] = round(num(ws.cell(r, c).value), 2)
        data[cli] = arr
    return data

def parse_anticuacion(ws):
    """Bloque Anticuacion: {cliente:{corr, b1_30}} sumando filas por cliente."""
    hdr = buscar_fila(ws, lambda w, r: str(w.cell(r, 3).value or "").strip() == "1-30"
                      and str(w.cell(r, 4).value or "").strip() == "Corriente")
    data = {}
    cli = None
    for r in range(hdr + 1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        if a.startswith("Suma total"):
            break
        if a:
            cli = a
        if cli is None:
            continue
        d = data.setdefault(cli, {"corr": 0.0, "b1_30": 0.0})
        d["b1_30"] += num(ws.cell(r, 3).value)   # col C = 1-30
        d["corr"]  += num(ws.cell(r, 4).value)   # col D = Corriente
    return {k: {"corr": round(v["corr"], 2), "b1_30": round(v["b1_30"], 2)}
            for k, v in data.items() if v["corr"] or v["b1_30"]}

def parse_pendientes(ws):
    """Tab 'Factura a cobrar': lista de facturas + proyeccion semanal."""
    pend, proj = [], []
    for r in range(2, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        if not a or a.startswith("Semana") or a.startswith("Total"):
            continue
        monto = ws.cell(r, 3).value
        if not isinstance(monto, (int, float)):
            continue
        if str(ws.cell(r, 2).value or "").strip() == "" and not isinstance(monto, (int, float)):
            continue
        pend.append({
            "cli":   a,
            "fc":    str(ws.cell(r, 2).value or "").strip(),
            "sin":   round(num(ws.cell(r, 3).value), 2),
            "iva":   round(num(ws.cell(r, 4).value), 2),
            "total": round(num(ws.cell(r, 5).value), 2),
            "venc":  ws.cell(r, 6).value if es_fecha(ws.cell(r, 6).value) else None,
            "est":   round(num(ws.cell(r, 7).value), 2),
            "prob":  ws.cell(r, 8).value if es_fecha(ws.cell(r, 8).value) else None,
            "pct":   round(num(ws.cell(r, 9).value) * 100, 2),
            "mail":  str(ws.cell(r, 10).value or "").strip(),
        })
    # proyeccion semanal (col A "Semana ...", col B valor sin FLK)
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        if a.startswith("Semana"):
            proj.append({"sem": a.replace("Semana", "").replace("de ", "").strip(),
                         "val": round(num(ws.cell(r, 2).value), 2)})
    return pend, proj


# ---------------------------------------------------------------- emision JS
def js_date(d):
    if not d or d.year < 2000:
        return "null"
    return f"new Date({d.year},{d.month - 1},{d.day})"

def js_dict_arr(d):
    return "{\n" + ",\n".join(f"  {json.dumps(k, ensure_ascii=False)}: {json.dumps(v)}"
                              for k, v in d.items()) + "\n}"

def js_pending(pend):
    out = []
    for p in pend:
        prob = "—" if not p["prob"] or p["prob"].year < 2000 else \
               f'{p["prob"].day:02d}/{p["prob"].month:02d}/{str(p["prob"].year)[2:]}'
        out.append("  {{cli:{cli}, fc:{fc}, sin:{sin}, iva:{iva}, total:{total}, "
                   "venc:{venc}, est:{est}, prob:{prob}, pct:{pct}, mail:{mail}}}".format(
                       cli=json.dumps(p["cli"], ensure_ascii=False),
                       fc=json.dumps(p["fc"], ensure_ascii=False),
                       sin=p["sin"], iva=p["iva"], total=p["total"],
                       venc=js_date(p["venc"]), est=p["est"],
                       prob=json.dumps(prob, ensure_ascii=False),
                       pct=p["pct"], mail=json.dumps(p["mail"], ensure_ascii=False)))
    return "[\n" + ",\n".join(out) + "\n]"

def js_antic(d):
    return "{\n" + ",\n".join(f"  {json.dumps(k, ensure_ascii=False)}: {{corr:{v['corr']}, b1_30:{v['b1_30']}}}"
                              for k, v in d.items()) + "\n}"


def construir_datos(meses, fact, cobr, antic, pend, proj, fecha):
    return f"""const REPORT_DATE = new Date({fecha.year},{fecha.month - 1},{fecha.day}); // {fecha:%d/%m/%Y}
const MESES = {json.dumps(meses, ensure_ascii=False)};
const FACT = {js_dict_arr(fact)};
const COBR = {js_dict_arr(cobr)};
const PENDING = {js_pending(pend)};
const ANTIC = {js_antic(antic)};
const PROJ = {json.dumps(proj, ensure_ascii=False)};"""


# ---------------------------------------------------------------- main
def obtener_xlsx(args):
    paths = [a for a in args if a.lower().endswith(".xlsx")]
    if paths:
        return paths[0]
    # intentar descarga directa (solo si la planilla es accesible por enlace)
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    tmp = "_planilla_descargada.xlsx"
    try:
        print("Descargando planilla desde Google Drive...")
        urllib.request.urlretrieve(url, tmp)
        return tmp
    except Exception as e:
        sys.exit("No se pudo descargar la planilla automaticamente (%s).\n"
                 "Descargala como .xlsx (Archivo > Descargar > Microsoft Excel) y pasala como argumento:\n"
                 "    python generar_dashboard.py archivo.xlsx" % e)


def main():
    args = sys.argv[1:]
    fecha = datetime.date.today()
    if "--fecha" in args:
        i = args.index("--fecha")
        fecha = datetime.datetime.strptime(args[i + 1], "%d/%m/%Y").date()
        del args[i:i + 2]

    xlsx = obtener_xlsx(args)
    if not os.path.exists(xlsx):
        sys.exit("No existe el archivo: " + xlsx)
    if not os.path.exists(PLANTILLA):
        sys.exit("Falta la plantilla '%s' (debe estar junto a este script)." % PLANTILLA)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    tb = wb["Tableros"]
    fc = wb["Factura a cobrar"]

    meses, fact = parse_facturacion(tb)
    cobr        = parse_ingresos(tb, meses)
    antic       = parse_anticuacion(tb)
    pend, proj  = parse_pendientes(fc)

    datos = construir_datos(meses, fact, cobr, antic, pend, proj, fecha)
    html = open(PLANTILLA, encoding="utf-8").read().replace("/*__DATOS_AUTOGENERADOS__*/", datos)

    salida = f"DSH_SEGUIMIENTO_CLIENTES_WS_{fecha:%d%m%Y}.html"
    open(salida, "w", encoding="utf-8").write(html)

    tot_f = sum(sum(v) for v in fact.values())
    tot_c = sum(sum(v) for v in cobr.values())
    print("OK ->", salida)
    print(f"  Clientes facturados : {len(fact)}")
    print(f"  Facturado 2026      : ${tot_f:,.0f}")
    print(f"  Cobrado 2026        : ${tot_c:,.0f}")
    print(f"  Facturas pendientes : {len(pend)}  (estimado ${sum(p['est'] for p in pend):,.0f})")


if __name__ == "__main__":
    main()
